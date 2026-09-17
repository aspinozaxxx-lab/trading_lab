"""Read-only metadata gate for saved CBR appendices; never return numeric magnitudes."""

import argparse
import hashlib
import json
import math
import re
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

MONTHS = [
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
]


def normalized(value):
    return " ".join(str(value or "").lower().replace("ё", "е").split())


def inspect_book(path, report_month):
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        matches = [
            s
            for s in book.worksheets
            if "чистые продажи иностранной валюты" in normalized(s["A1"].value)
            and "экспортер" in normalized(s["A1"].value)
        ]
        if not matches:
            return {"status": "NO_NET_SALES_SERIES", "sheets": len(book.worksheets)}
        if len(matches) != 1:
            return {"status": "AMBIGUOUS_NET_SALES_SERIES"}
        sheet = matches[0]
        header = [
            c
            for row in sheet.iter_rows(min_row=1, max_row=8)
            for c in row
            if isinstance(c.value, str)
        ]
        labels = {c.coordinate: c.value for c in header}
        result = {"sheet": sheet.title, "labels": labels}
        text = " ".join(normalized(c.value) for c in header)
        result["cohort_29_labeled"] = bool(re.search(r"\b29\b", text))
        result["billion_usd_labeled"] = "млрд долл. сша" in text
        columns = {normalized(c.value): c for c in header}
        value_header = columns.get("чистые продажи")
        date_header = columns.get("месяц")
        month_header, year_header = columns.get("месяцы"), columns.get("годы")
        if value_header is None or not (date_header or (month_header and year_header)):
            return {**result, "status": "UNSUPPORTED_LAYOUT"}
        date_headers = [date_header] if date_header else [month_header, year_header]
        if any(c.row != value_header.row for c in date_headers):
            return {**result, "status": "MISALIGNED_HEADERS"}
        rows, year = {}, None
        width = max(c.column for c in date_headers + [value_header])
        for cells in sheet.iter_rows(min_row=value_header.row + 1, max_col=width):
            number = next((c.row for c in cells if c.value is not None), None)
            if number is None:
                continue
            if date_header:
                day = cells[date_header.column - 1].value
                if day is None:
                    continue
                if not isinstance(day, datetime) or day.day != 1:
                    return {**result, "status": "UNSUPPORTED_DATE_CELL", "row": number}
                period = day.strftime("%Y-%m")
            else:
                year_value = cells[year_header.column - 1].value
                month = normalized(cells[month_header.column - 1].value)
                if year_value is None and not month:
                    continue
                if year_value is not None:
                    if not isinstance(year_value, int) or isinstance(year_value, bool):
                        return {**result, "status": "UNSUPPORTED_YEAR_CELL", "row": number}
                    year = year_value
                if year is None or month not in MONTHS:
                    return {**result, "status": "UNSUPPORTED_MONTH_CELL", "row": number}
                period = f"{year}-{MONTHS.index(month) + 1:02}"
            if period >= "2026-01":
                return {**result, "status": "PROTECTED_DATE_REJECTED", "row": number}
            if period in rows:
                return {**result, "status": "DUPLICATE_MONTH", "period": period}
            cell = cells[value_header.column - 1]
            # Presence/type only. Do not print/derive magnitudes, signs or market series.
            valid = (
                cell.data_type == "n"
                and isinstance(cell.value, (int, float))
                and not isinstance(cell.value, bool)
                and math.isfinite(cell.value)
            )
            rows[period] = {
                "date_cells": [f"{c.column_letter}{number}" for c in date_headers],
                "value_cell": f"{value_header.column_letter}{number}",
                "numeric_cell_present": valid,
            }
        y, m = map(int, report_month.split("-"))
        previous = f"{y if m > 1 else y - 1}-{m - 1 if m > 1 else 12:02}"
        complete = all(
            rows.get(p, {}).get("numeric_cell_present", False) for p in (previous, report_month)
        )
        ready = complete and result["cohort_29_labeled"] and result["billion_usd_labeled"]
        return {
            **result,
            "status": "READY_PAIR_METADATA" if ready else "NOT_ADMITTED_PAIR",
            "dates_and_presence_only": rows,
            "previous_month": previous,
            "report_month": report_month,
            "pair_present": complete,
        }
    finally:
        book.close()


def main():
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("root", type=Path)
    cli.add_argument("--output", required=True, type=Path)
    args = cli.parse_args()
    manifest = json.loads((args.root / "manifest.json").read_text(encoding="utf-8-sig"))
    assert manifest["status"] == "COMPLETE_RAW_ONLY"
    for name, digest in manifest["files"].items():
        assert hashlib.sha256((args.root / name).read_bytes()).hexdigest() == digest, name
    records = [
        {"report_month": period, **inspect_book(args.root / (period + ".xlsx"), period)}
        for period in manifest["completed_periods"]
    ]
    result = {
        "status": "COMPLETE_METADATA_ONLY",
        "economic_admission": False,
        "completed_at_utc": datetime.now(UTC).isoformat(),
        "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "source_manifest_sha256": hashlib.sha256(
            (args.root / "manifest.json").read_bytes()
        ).hexdigest(),
        "counts": dict(Counter(r["status"] for r in records)),
        "records": records,
    }
    with args.output.open("x", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(
        json.dumps(
            {
                "counts": result["counts"],
                "periods": {r["report_month"]: r["status"] for r in records},
                "report_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
            },
            ensure_ascii=False,
        ),
        flush=True,
    )


if __name__ == "__main__":
    main()
