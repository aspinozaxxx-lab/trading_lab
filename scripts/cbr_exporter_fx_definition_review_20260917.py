"""Separate daily/monthly CBR source layouts without aggregating or exposing values."""

import argparse
import hashlib
import json
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import cbr_exporter_fx_metadata_20260917 as metadata
from openpyxl import load_workbook


def review(root, record):
    period = record["report_month"]
    if record["status"] != "UNSUPPORTED_LAYOUT":
        return {"report_month": period, "original_status": record["status"]}
    book = load_workbook(
        root / (period + ".xlsx"), read_only=True, data_only=False, keep_links=False
    )
    try:
        sheet = book[record["sheet"]]
        headers = [
            c
            for row in sheet.iter_rows(min_row=1, max_row=8)
            for c in row
            if metadata.normalized(c.value) == "чистые продажи"
        ]
        assert len(headers) == 1
        header = headers[0]
        date_column = header.column - 1
        assert date_column >= 1
        dates = []
        for row in sheet.iter_rows(min_row=header.row + 1, max_col=date_column):
            value = row[date_column - 1].value
            if value is None:
                continue
            assert isinstance(value, datetime), (period, "unrecognized date column")
            assert value.year <= 2025
            dates.append(value.date().isoformat())
        assert dates and dates == sorted(set(dates))
        monthly = all(d.endswith("-01") for d in dates)
        return {
            "report_month": period,
            "sheet": sheet.title,
            "original_status": record["status"],
            "date_grain": "monthly_first_day" if monthly else "daily_dates",
            "rows_with_date": len(dates),
            "first_date": dates[0],
            "last_date": dates[-1],
            "date_column": date_column,
            "value_header": header.coordinate,
            "labels": record["labels"],
            "magnitude_or_aggregation_performed": False,
        }
    finally:
        book.close()


def main():
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("root", type=Path)
    cli.add_argument("metadata", type=Path)
    cli.add_argument("--output", required=True, type=Path)
    args = cli.parse_args()
    previous = json.loads(args.metadata.read_text(encoding="utf-8"))
    assert hashlib.sha256(args.metadata.read_bytes()).hexdigest() == (
        "e44ca3a2e598737ff98a57212ef210ab9de6b0d1de4cfc6ec5b12e1e304db7a9"
    )
    assert (
        hashlib.sha256((args.root / "manifest.json").read_bytes()).hexdigest()
        == (previous["source_manifest_sha256"])
    )
    manifest = json.loads((args.root / "manifest.json").read_text(encoding="utf-8-sig"))
    for name, digest in manifest["files"].items():
        assert hashlib.sha256((args.root / name).read_bytes()).hexdigest() == digest
    records = [review(args.root, row) for row in previous["records"]]
    result = {
        "status": "COMPLETE_DEFINITION_REVIEW_ONLY",
        "economic_admission": False,
        "completed_at_utc": datetime.now(UTC).isoformat(),
        "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "prior_metadata_sha256": hashlib.sha256(args.metadata.read_bytes()).hexdigest(),
        "counts": dict(Counter(r.get("date_grain", r["original_status"]) for r in records)),
        "records": records,
    }
    with args.output.open("x", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(
        json.dumps(
            {
                "counts": result["counts"],
                "layouts": [
                    {
                        k: r[k]
                        for k in (
                            "report_month",
                            "date_grain",
                            "rows_with_date",
                            "first_date",
                            "last_date",
                        )
                    }
                    for r in records
                    if "date_grain" in r
                ],
                "report_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
